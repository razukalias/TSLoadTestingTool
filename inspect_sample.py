from openpyxl import load_workbook
from pathlib import Path
p = Path(__file__).parent / 'samples/jsonplaceholder-rest/DataEngine.xlsx'
wb = load_workbook(p, read_only=True, data_only=False)
for name in ['config','request','response']:
    ws = wb[name]
    print(f'[{name}]')
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        while vals and vals[-1] is None:
            vals.pop()
        print(vals)
