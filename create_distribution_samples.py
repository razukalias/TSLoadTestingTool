from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import json
import shutil

ROOT = Path('/home/ubuntu/LoadTestingTool-distribution')
INSTANCES = ROOT / 'instances'


def workbook(path, config_rows, request_headers, request_rows, response_headers, response_rows, appsettings, templates=None, extra_files=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    appsettings.setdefault('InternalLoggingEnabled', True)
    templates_dir = path.parent / 'Templates'
    templates_dir.mkdir(exist_ok=True)
    if templates:
        for name, content in templates.items():
            (templates_dir / name).write_text(content, encoding='utf-8')
    for directory in ['Results', 'History', 'Logs', 'Inputs', 'Queries']:
        (path.parent / directory).mkdir(exist_ok=True)
    if extra_files:
        for name, content in extra_files.items():
            target = path.parent / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding='utf-8')
    wb = Workbook()
    config = wb.active
    config.title = 'config'
    config.append(list(config_rows[0].keys()))
    for row in config_rows:
        config.append([row[h] for h in config_rows[0].keys()])
    req = wb.create_sheet('request')
    req.append(request_headers)
    for row in request_rows:
        req.append(row)
    resp = wb.create_sheet('response')
    resp.append(response_headers)
    for row in response_rows:
        resp.append(row)
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9EAF7')
        for column in sheet.columns:
            letter = column[0].column_letter
            values = [len(str(c.value or '')) for c in column]
            sheet.column_dimensions[letter].width = min(55, max(14, max(values, default=14) + 2))
    wb.save(path)
    (path.parent / 'appsettings.json').write_text(json.dumps(appsettings, indent=2) + '\n', encoding='utf-8')


def main():
    if ROOT.exists():
        shutil.rmtree(ROOT)
    INSTANCES.mkdir(parents=True)

    # Online HTTP: GET, POST, assertions, correlation, data IDs, and environments.
    http = INSTANCES / 'DataEngine_OnlineHttp'
    workbook(
        http / 'DataEngine.xlsx',
        [
            {'testcaseindex': 1, 'testcase': 'JsonPlaceholderFlow', 'stepname': 'FetchPost', 'steptype': 'HTTP', 'templetsource': 'empty.json', 'targetUrl': 'https://jsonplaceholder.typicode.com/posts/1', 'verb': 'GET', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '200', 'stoponfailure': True, 'environments': 'online', 'header_Accept': 'application/json'},
            {'testcaseindex': 1, 'testcase': 'JsonPlaceholderFlow', 'stepname': 'CreatePost', 'steptype': 'HTTP', 'templetsource': 'create-post.json', 'targetUrl': 'https://jsonplaceholder.typicode.com/posts', 'verb': 'POST', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 100, 'expectedstatus': '201', 'stoponfailure': True, 'environments': 'online', 'header_Accept': 'application/json', 'header_Content-Type': 'application/json'},
        ],
        ['testcaseindex', 'testcase', 'dataid', 'CreatePost.title', 'CreatePost.body'],
        [[1, 'JsonPlaceholderFlow', 'Online001', 'First online sample', 'Created by LoadTestingTool.'], [1, 'JsonPlaceholderFlow', 'Online002', 'Second online sample', 'A second isolated data row.']],
        ['testcaseindex', 'testcase', 'dataid', 'FetchPost.statusCode', 'FetchPost.id', 'FetchPost.title', 'FetchPost.userId', 'FetchPost.body', 'CreatePost.statusCode', 'CreatePost.id', 'CreatePost.title', 'CreatePost.body', 'CreatePost.userId', 'CreatePost.sourcePostId', 'extractvariable'],
        [[1, 'JsonPlaceholderFlow', 'Online001', 'eq:200', 'eq:1', 'notempty', 'eq:1', 'notempty', 'eq:201', 'eq:101', 'eq:First online sample', 'eq:Created by LoadTestingTool.', 'eq:1', 'eq:1', 'FetchPost.id:postId'], [1, 'JsonPlaceholderFlow', 'Online002', 'eq:200', 'eq:1', 'notempty', 'eq:1', 'notempty', 'eq:201', 'eq:101', 'eq:Second online sample', 'eq:A second isolated data row.', 'eq:1', 'eq:1', 'FetchPost.id:postId']],
        {'ExcelFilePath': 'DataEngine.xlsx', 'TemplatesFolder': 'Templates', 'ResultsFolder': 'Results', 'HistoryFolder': 'History', 'LogsFolder': 'Logs', 'RequestTimeoutSeconds': 30, 'TestcaseSelection': '0', 'EnvironmentSelection': 'online', 'MaskSensitiveData': True},
        {'empty.json': '{}\n', 'create-post.json': '{\n  "title": "_title_",\n  "body": "_body_",\n  "userId": 1,\n  "sourcePostId": "_postId_"\n}\n'},
    )
    (http / 'README.md').write_text('# Online HTTP sample\n\nRuns GET /posts/1 followed by POST /posts against JSONPlaceholder. It demonstrates two isolated data IDs, JSON assertions, response correlation, headers, wait timing, and the online environment.\n\nRun with `--testcases 0`. The POST endpoint returns a synthetic response and does not permanently save data.\n', encoding='utf-8')

    # Online GraphQL.
    gql = INSTANCES / 'DataEngine_OnlineGraphQL'
    gql_config = json.dumps({'operationName': 'Country', 'variables': {'code': '_countryCode_'}}, separators=(',', ':'))
    workbook(
        gql / 'DataEngine.xlsx',
        [{'testcaseindex': 1, 'testcase': 'GraphQLCountry', 'stepname': 'LookupCountry', 'steptype': 'GraphQL', 'templetsource': 'country.graphql', 'targetUrl': 'https://countries.trevorblades.com/', 'verb': 'POST', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '200', 'stoponfailure': True, 'environments': 'online', 'stepconfig': gql_config}],
        ['testcaseindex', 'testcase', 'dataid', 'LookupCountry.countryCode'],
        [[1, 'GraphQLCountry', 'GraphQL001', 'US']],
        ['testcaseindex', 'testcase', 'dataid', 'LookupCountry.statusCode', 'LookupCountry.data.country.code', 'LookupCountry.data.country.name', 'LookupCountry.data.country.capital'],
        [[1, 'GraphQLCountry', 'GraphQL001', 'eq:200', 'eq:US', 'eq:United States', 'notempty']],
        {'ExcelFilePath': 'DataEngine.xlsx', 'TemplatesFolder': 'Templates', 'ResultsFolder': 'Results', 'HistoryFolder': 'History', 'LogsFolder': 'Logs', 'RequestTimeoutSeconds': 30, 'TestcaseSelection': '0', 'EnvironmentSelection': 'online', 'MaskSensitiveData': True},
    )
    (gql / 'Templates' / 'country.graphql').write_text('query Country($code: ID!) { country(code: $code) { code name capital } }\n', encoding='utf-8')
    (gql / 'README.md').write_text('# Online GraphQL sample\n\nQueries the public Countries GraphQL API for the United States. It demonstrates a GraphQL template, operation metadata, variable substitution, and nested response assertions.\n', encoding='utf-8')

    # File and trusted script.
    local = INSTANCES / 'DataEngine_FileAndScript'
    file_config = json.dumps({'outputVariable': 'customerData'}, separators=(',', ':'))
    script_config = '{}'
    workbook(
        local / 'DataEngine.xlsx',
        [
            {'testcaseindex': 1, 'testcase': 'LocalFeatures', 'stepname': 'ReadCustomer', 'steptype': 'File', 'templetsource': 'read-customer.json', 'targetUrl': '', 'verb': 'POST', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '', 'stoponfailure': True, 'environments': 'local', 'stepconfig': file_config},
            {'testcaseindex': 1, 'testcase': 'LocalFeatures', 'stepname': 'NormalizeName', 'steptype': 'Script', 'templetsource': 'normalize-name.csx', 'targetUrl': '', 'verb': 'POST', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '', 'stoponfailure': True, 'environments': 'local', 'stepconfig': script_config},
        ],
        ['testcaseindex', 'testcase', 'dataid'], [[1, 'LocalFeatures', 'Local001']],
        ['testcaseindex', 'testcase', 'dataid', 'ReadCustomer.statusCode', 'ReadCustomer.customer.id', 'ReadCustomer.customer.name'], [[1, 'LocalFeatures', 'Local001', 'eq:200', 'eq:7', 'eq:Alice']],
        {'ExcelFilePath': 'DataEngine.xlsx', 'TemplatesFolder': 'Templates', 'ResultsFolder': 'Results', 'HistoryFolder': 'History', 'LogsFolder': 'Logs', 'RequestTimeoutSeconds': 30, 'TestcaseSelection': '0', 'EnvironmentSelection': 'local', 'WorkspaceRoot': '.', 'AllowTrustedScripts': True, 'ScriptTimeoutSeconds': 30, 'MaskSensitiveData': True},
        templates={'read-customer.json': '{"action":"JSON","path":"Inputs/customer.json"}\n', 'normalize-name.csx': 'Runtime.Set("normalizedName", Runtime.Get("customerData")?.ToString()?.ToUpperInvariant()); Runtime.Get("normalizedName")\n'},
        extra_files={'Inputs/customer.json': '{\n  "customer": {"id": 7, "name": "Alice"}\n}\n'},
    )
    (local / 'README.md').write_text('# Local file and script sample\n\nReads `Inputs/customer.json` with a File step and runs a trusted C# Script step. It does not require a network or database. `AllowTrustedScripts` is enabled only for this controlled sample.\n', encoding='utf-8')

    # SQL Server sample for local Docker SQL Server.
    sql = INSTANCES / 'DataEngine_SqlServer'
    sql_config = json.dumps({'connectionProfile': 'LocalSql', 'commandType': 'Text', 'parameters': {'id': '_customerId_'}, 'outputs': {'customerId': 'Id', 'customerName': 'Name', 'customerStatus': 'Status'}}, separators=(',', ':'))
    workbook(
        sql / 'DataEngine.xlsx',
        [{'testcaseindex': 1, 'testcase': 'SqlSmokeTest', 'stepname': 'GetCustomer', 'steptype': 'SQL', 'templetsource': 'get-customer.sql', 'targetUrl': '', 'verb': 'POST', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '', 'stoponfailure': True, 'environments': 'local', 'stepconfig': sql_config}],
        ['testcaseindex', 'testcase', 'dataid', 'GetCustomer.customerId'], [[1, 'SqlSmokeTest', 'Sql001', '1']],
        ['testcaseindex', 'testcase', 'dataid', 'GetCustomer.[0].Id', 'GetCustomer.[0].Name', 'GetCustomer.[0].Status'], [[1, 'SqlSmokeTest', 'Sql001', 'eq:1', 'eq:Alice', 'eq:ACTIVE']],
        {'ExcelFilePath': 'DataEngine.xlsx', 'TemplatesFolder': 'Templates', 'ResultsFolder': 'Results', 'HistoryFolder': 'History', 'LogsFolder': 'Logs', 'RequestTimeoutSeconds': 30, 'TestcaseSelection': '0', 'EnvironmentSelection': 'local', 'WorkspaceRoot': '.', 'MaskSensitiveData': True, 'Connections': {'LocalSql': {'ConnectionString': 'Server=localhost,1433;Database=LoadTest;User Id=sa;Password=LoadTesting#Pass123;Encrypt=False;TrustServerCertificate=True', 'UseIntegratedSecurity': False}}},
        templates={'get-customer.sql': 'SELECT Id, Name, Status FROM Customers WHERE Id = @id;\n'},
    )
    (sql / 'README.md').write_text('# SQL Server sample\n\nRequires the local Docker SQL Server from the project instructions. Create database `LoadTest`, table `Customers`, and row `(1, "Alice", "ACTIVE")`; then run this instance.\n', encoding='utf-8')

    # Online failure/diagnostics sample.
    fail = INSTANCES / 'DataEngine_OnlineFailure'
    workbook(
        fail / 'DataEngine.xlsx',
        [{'testcaseindex': 1, 'testcase': 'FailureDiagnostics', 'stepname': 'ExpectedFailure', 'steptype': 'HTTP', 'templetsource': 'empty.json', 'targetUrl': 'https://httpbin.org/status/418', 'verb': 'GET', 'contenttype': 'json', 'threads': 1, 'iterations': 1, 'wait': 0, 'expectedstatus': '200', 'stoponfailure': True, 'environments': 'online'}],
        ['testcaseindex', 'testcase', 'dataid'], [[1, 'FailureDiagnostics', 'Fail001']],
        ['testcaseindex', 'testcase', 'dataid', 'ExpectedFailure.statusCode'], [[1, 'FailureDiagnostics', 'Fail001', 'eq:418']],
        {'ExcelFilePath': 'DataEngine.xlsx', 'TemplatesFolder': 'Templates', 'ResultsFolder': 'Results', 'HistoryFolder': 'History', 'LogsFolder': 'Logs', 'RequestTimeoutSeconds': 30, 'TestcaseSelection': '0', 'EnvironmentSelection': 'online', 'MaskSensitiveData': True},
        {'empty.json': '{}\n'},
    )
    (fail / 'README.md').write_text('# Online failure sample\n\nCalls httpbin `/status/418` while expecting status 200. It should fail and demonstrates failed-result folders, error history, and exit code 2.\n', encoding='utf-8')

    (ROOT / 'README.md').write_text('''# LoadTestingTool Windows distribution\n\n## Folders\n\n- `runner`: Release console runner output. Run it with `dotnet runner/LoadTestingTool.dll ...`.\n- `ui`: Release Avalonia UI output. It is configured to discover `instances/DataEngine_*`.\n- `instances`: Ready-made workbooks and templates.\n\n## Online tests\n\n`DataEngine_OnlineHttp` uses JSONPlaceholder. `DataEngine_OnlineGraphQL` uses the public Countries GraphQL API. `DataEngine_OnlineFailure` intentionally fails against httpbin status 418.\n\n## Local tests\n\n`DataEngine_FileAndScript` is self-contained. `DataEngine_SqlServer` requires a local Docker SQL Server and the setup described in its README.\n\n## Console example\n\n```powershell\ndotnet runner\\LoadTestingTool.dll --excel instances\\DataEngine_OnlineHttp\\DataEngine.xlsx --templates instances\\DataEngine_OnlineHttp\\Templates --results instances\\DataEngine_OnlineHttp\\Results --testcases 0\n```\n\nThe UI requires the .NET 10 Desktop Runtime on Windows. Feature-step operation content is stored in each instance's `Templates` folder; `stepconfig` contains metadata such as connection profiles, parameters, paths, and output names.\n''', encoding='utf-8')
    print(ROOT)


if __name__ == '__main__':
    main()
