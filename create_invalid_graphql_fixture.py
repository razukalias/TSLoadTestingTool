from pathlib import Path
from openpyxl import load_workbook
import shutil
source = Path('/home/ubuntu/LoadTestingTool-distribution/instances/DataEngine_OnlineGraphQL/DataEngine.xlsx')
target_root = Path('/home/ubuntu/log-validation-fixture')
if target_root.exists(): shutil.rmtree(target_root)
target_root.mkdir(parents=True)
shutil.copytree(source.parent, target_root / 'instance')
path = target_root / 'instance' / 'DataEngine.xlsx'
wb = load_workbook(path)
ws = wb['config']
headers = {str(c.value).lower(): c.column for c in ws[1]}
ws.cell(2, headers['stepconfig']).value = ''
wb.save(path)
print(path)
