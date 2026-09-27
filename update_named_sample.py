from openpyxl import load_workbook
p='/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder/DataEngine.xlsx'
wb=load_workbook(p)
ws=wb['response']
headers=['testcaseindex','testcase','FetchPost.statusCode','FetchPost.userId','FetchPost.id','FetchPost.title','FetchPost.body','CreatePost.statusCode','CreatePost.id','extractvariable']
values=[1,'GetThenPost','eq:200','eq:1','eq:1','notempty','notempty','eq:201','eq:101','FetchPost.id:postId']
for col in range(1, ws.max_column + 1):
    ws.cell(1,col).value=None
    ws.cell(2,col).value=None
for col, value in enumerate(headers,1): ws.cell(1,col).value=value
for col, value in enumerate(values,1): ws.cell(2,col).value=value
wb.save(p)
