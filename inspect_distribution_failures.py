from pathlib import Path
from openpyxl import load_workbook
import sys
root = Path(sys.argv[1])
for p in sorted(root.rglob('*.xlsx')):
    print(f'FILE {p}')
    wb = load_workbook(p, read_only=True, data_only=True)
    if 'AssertionHistory' in wb.sheetnames:
        ws = wb['AssertionHistory']
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[1:]:
            if row[10] == 'FAIL':
                print('FAIL', row[4], row[5], row[6], row[7], 'expected=', row[8], 'actual=', row[9], 'message=', row[11])
