from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).parent
for path in [root / 'samples/jsonplaceholder-rest/DataEngine.xlsx', root / 'Requirements_Questionnaire.xlsx']:
    print(f'FILE: {path.name}')
    wb = load_workbook(path, read_only=True, data_only=False)
    print('SHEETS:', wb.sheetnames)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        print(f' SHEET {ws.title}: headers={header}')
        for i, row in zip(range(1, 8), rows):
            print('  ', row)
    print()
