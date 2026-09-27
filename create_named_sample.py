from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import json

root = Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder')
templates = root / 'Templates'
templates.mkdir(parents=True, exist_ok=True)
(root / 'History').mkdir(exist_ok=True)
(root / 'Logs').mkdir(exist_ok=True)
(root / 'Results').mkdir(exist_ok=True)

(root / 'Templates' / 'fetch-post.json').write_text('{}\n', encoding='utf-8')
(root / 'Templates' / 'create-post.json').write_text(json.dumps({
    'title': '_title_',
    'body': '_body_',
    'userId': 1,
    'sourcePostId': '_postId_'
}, indent=2) + '\n', encoding='utf-8')

wb = Workbook()
ws = wb.active
ws.title = 'config'
config_headers = ['testcaseindex', 'testcase', 'stepname', 'templetsource', 'targetUrl', 'verb', 'contenttype', 'threads', 'iterations', 'wait', 'expectedstatus', 'stoponfailure', 'environments', 'header_Accept', 'header_Content-Type']
ws.append(config_headers)
ws.append([1, 'GetThenPost', 'FetchPost', 'fetch-post.json', 'https://jsonplaceholder.typicode.com/posts/1', 'GET', 'json', 1, 1, 0, '200', True, 'online', 'application/json', ''])
ws.append([1, 'GetThenPost', 'CreatePost', 'create-post.json', 'https://jsonplaceholder.typicode.com/posts', 'POST', 'json', 1, 1, 250, '201', True, 'online', 'application/json', 'application/json'])

req = wb.create_sheet('request')
req.append(['testcaseindex', 'testcase', 'dataid', 'CreatePost.title', 'CreatePost.body'])
req.append([1, 'GetThenPost', 'Online001', 'Named-step payload demo', 'Created by the named-step Excel sample.'])

resp = wb.create_sheet('response')
resp.append(['testcaseindex', 'testcase', 'FetchPost.statusCode', 'FetchPost.id', 'FetchPost.title', 'CreatePost.statusCode', 'CreatePost.id', 'CreatePost.sourcePostId', 'extractvariable'])
resp.append([1, 'GetThenPost', 'eq:200', 'eq:1', 'notempty', 'eq:201', 'eq:101', 'eq:1', 'FetchPost.id:postId'])

for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(45, max(14, max(len(str(c.value or '')) for c in column) + 2))

wb.save(root / 'DataEngine.xlsx')
(root / 'appsettings.json').write_text(json.dumps({
    'ExcelFilePath': 'DataEngine.xlsx',
    'TemplatesFolder': 'Templates',
    'LogsFolder': 'Logs',
    'HistoryFolder': 'History',
    'ResultsFolder': 'Results',
    'RequestTimeoutSeconds': 30,
    'RunScenariosInParallel': False,
    'TestcaseSelection': '0',
    'PromptForTestcaseSelection': False,
    'MaskSensitiveData': True
}, indent=2) + '\n', encoding='utf-8')
(root / 'README.md').write_text('''# Named-step online GET/POST sample\n\nThis standalone sample is separate from the original archive and demonstrates the redesigned workbook model. It uses the public JSONPlaceholder REST API and requires no credentials.\n\n## Flow\n\n`GetThenPost` contains two named request steps:\n\n1. `FetchPost` performs `GET /posts/1`, asserts the response, and extracts `FetchPost.id` as `postId`.\n2. `CreatePost` performs `POST /posts`, uses `_postId_` in the JSON payload, and asserts the created response.\n\n## Workbook model\n\n- `config`: one row per named HTTP request step; row order defines execution order.\n- `request`: one row for the testcase; variables use `stepname.variable` headers.\n- `response`: one row for the testcase; assertion headers use `stepname.responsepath`.\n- `extractvariable`: pipe-separated mappings such as `FetchPost.id:postId|CreatePost.id:createdId`.\n\n## Run\n\nFrom this directory, after building the console runner:\n\n```bash\ndotnet path/to/LoadTestingTool.dll --excel DataEngine.xlsx --templates Templates --history History --logs Logs --results Results --testcases 0\n```\n\nThe POST endpoint is a public demonstration endpoint. It normally returns HTTP 201 and a synthetic ID; it does not persist data permanently.\n''', encoding='utf-8')
print(root)
