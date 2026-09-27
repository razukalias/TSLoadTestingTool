from pathlib import Path
from openpyxl import load_workbook
root = Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder')
for p in sorted((root/'History').glob('*.xlsx')):
    print('FILE', p)
    wb = load_workbook(p, read_only=True, data_only=True)
    for name in ['RunSummary','RequestHistory','AssertionHistory','CorrelationHistory']:
        ws = wb[name]
        print('SHEET', name)
        for row in ws.iter_rows(values_only=True): print(list(row))
for p in sorted((root/'Logs').rglob('*.log')):
    print('LOG', p)
    print(p.read_text(errors='replace')[-6000:])
