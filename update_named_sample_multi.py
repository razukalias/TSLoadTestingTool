from openpyxl import load_workbook
from pathlib import Path
p=Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder/DataEngine.xlsx')
wb=load_workbook(p)
req=wb['request']
for row in list(req.iter_rows(min_row=2, max_row=req.max_row)): req.delete_rows(row[0].row, 1)
req.append([1,'GetThenPost','Online001','Named-step payload demo A','Created by data row A.'])
req.append([1,'GetThenPost','Online002','Named-step payload demo B','Created by data row B.'])
resp=wb['response']
for row in list(resp.iter_rows(min_row=2, max_row=resp.max_row)): resp.delete_rows(row[0].row, 1)
headers=['testcaseindex','testcase','dataid','FetchPost.statusCode','FetchPost.userId','FetchPost.id','FetchPost.title','FetchPost.body','CreatePost.statusCode','CreatePost.id','extractvariable']
for col in range(1, resp.max_column+1): resp.cell(1,col).value=None
for col,v in enumerate(headers,1): resp.cell(1,col).value=v
for dataid in ['Online001','Online002']:
    resp.append([1,'GetThenPost',dataid,'eq:200','eq:1','eq:1','notempty','notempty','eq:201','eq:101','FetchPost.id:postId'])
wb.save(p)
