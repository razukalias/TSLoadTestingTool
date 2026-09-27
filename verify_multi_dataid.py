from pathlib import Path
from openpyxl import load_workbook
root=Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder')
p=sorted((root/'History').glob('*.xlsx'))[-1]
wb=load_workbook(p, read_only=True, data_only=True)
for sheet_name in ['TestcaseSummary','RequestHistory','AssertionHistory','CorrelationHistory']:
    ws=wb[sheet_name]; rows=list(ws.iter_rows(values_only=True)); print(sheet_name, 'rows=', len(rows)-1)
    print('headers=', rows[0])
    for row in rows[1:]: print(row[:8])
