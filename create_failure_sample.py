from pathlib import Path
from shutil import copytree, copy2
from openpyxl import load_workbook
src=Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder')
dst=Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder-failure')
if dst.exists():
    import shutil; shutil.rmtree(dst)
copytree(src,dst,ignore=lambda d,n: {'History','Logs','Results'} if d==str(src) else set())
for name in ['History','Logs','Results']:
    (dst/name).mkdir()
wb=load_workbook(dst/'DataEngine.xlsx')
ws=wb['response']
# Force Online001 FetchPost.statusCode to fail while leaving Online002 valid.
for row in range(2, ws.max_row+1):
    if ws.cell(row,3).value == 'Online001':
        ws.cell(row,4).value='eq:999'
wb.save(dst/'DataEngine.xlsx')
