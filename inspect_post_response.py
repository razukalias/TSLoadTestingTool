from pathlib import Path
from openpyxl import load_workbook
root=Path('/home/ubuntu/work/loadtestingtool-analysis/samples/named-steps-jsonplaceholder')
p=sorted((root/'History').glob('*.xlsx'))[-1]
wb=load_workbook(p, read_only=True, data_only=True)
ws=wb['RequestHistory']
rows=list(ws.iter_rows(values_only=True))
for i,h in enumerate(rows[0]): print(i,h)
for row in rows[1:]:
    if row[4] == 'CreatePost':
        print('step', row[4], 'status', row[17])
        print('response body:', row[19])
